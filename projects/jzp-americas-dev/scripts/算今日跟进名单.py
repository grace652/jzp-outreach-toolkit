#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""算出「今天真正需要跟进」的客户（**以实测发出记录为准，不信表格**）

判据（关键）：
    实测最近发出日 L  vs  下次跟进日 D
      L <  D  → **逾期未跟**（今天要发）
      L >= D  → 已跟进（表格没推进，只需修表，不必重发）

为什么必须用实测：跟进表存在漂移 —— 9/29 批次实际发过 21 封，但表里 9/29 到期的有 26 家，
部分行没被推进。若按表格写，会对**已经跟过**的客户重复发。
"""

import importlib.util
import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

ROOT = Path(Path(__file__).resolve().parents[1])
XLSX = ROOT / "开发信项目/工作区/开发信跟进表.xlsx"
import os
TODAY = os.environ.get("JZP_TODAY", "2026-10-02")   # 可用 JZP_TODAY 覆盖
OUT = ROOT / f".workbuddy-ai/今日跟进名单-算-{TODAY}.json"
TERMINAL = {"已成交", "已退订（对方拒联）", "已放弃", "已退信", "个人处理", "无回应"}
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
MONTH = {m: i for i, m in enumerate(
    "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split(), 1)}
DATE_RE = re.compile(r"(\d{1,2})\s+([A-Z][a-z]{2})\s+(\d{4})")


def to_iso(s):
    m = DATE_RE.search(s or "")
    if not m:
        return None
    return f"{int(m.group(3)):04d}-{MONTH[m.group(2)]:02d}-{int(m.group(1)):02d}"


def load_mailer():
    spec = importlib.util.spec_from_file_location("mailer", ROOT / "邮箱客户端.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules["mailer"] = mod
    spec.loader.exec_module(mod)
    return mod


def main():
    wb = openpyxl.load_workbook(XLSX)
    rows = []
    for ws in wb.worksheets:
        for r in range(5, ws.max_row + 1):
            n = ws.cell(r, 1).value
            if n is None:
                continue
            nxt = str(ws.cell(r, 15).value or "")
            st = str(ws.cell(r, 13).value or "")
            if nxt in ("", "None", "—") or nxt > TODAY or st in TERMINAL:
                continue
            raw = str(ws.cell(r, 12).value or "")
            rows.append(dict(num=n, sheet=ws.title, name=str(ws.cell(r, 3).value or ""),
                             region=str(ws.cell(r, 4).value or ""), status=st,
                             cnt=ws.cell(r, 14).value, nxt=nxt,
                             last=str(ws.cell(r, 16).value or ""), raw=raw,
                             emails=[e.lower() for e in EMAIL_RE.findall(raw)]))

    mc = load_mailer()
    m = mc.imap_connect(mc.load_config())
    m.select('"Sent Messages"', readonly=True)
    typ, d = m.search(None, "ALL")
    per_dom = defaultdict(list)
    for num in d[0].split():
        raw = mc.fetch_headers(m, num, "TO CC SUBJECT DATE")
        blob = mc.raw_header(raw, "To") + " " + mc.raw_header(raw, "CC")
        iso = to_iso(mc.raw_header(raw, "Date"))
        subj = mc.raw_header(raw, "Subject")
        for e in set(x.lower() for x in EMAIL_RE.findall(blob)):
            per_dom[e].append((iso, e, subj))   # 🔴 按**精确地址**索引
    m.logout()

    for x in rows:
        seen, hits = set(), []
        for e in x["emails"]:
            for iso, ad, sj in per_dom.get(e, []):   # 🔴 精确地址匹配
                k = (ad, sj)
                if k in seen:
                    continue
                seen.add(k)
                hits.append((iso, ad, sj))
        hits = [h for h in hits if h[0]]
        hits.sort()
        x["sent_last"] = hits[-1][0] if hits else None
        x["sent_n"] = len(hits)
        x["sent_re"] = sum(1 for _, _, sj in hits if sj.lower().startswith("re:"))
        x["overdue"] = (x["sent_last"] is None) or (x["sent_last"] < x["nxt"])

    todo = [x for x in rows if x["overdue"]]
    done = [x for x in rows if not x["overdue"]]

    print(f"到期/逾期候选 {len(rows)} 家  →  🔴 真需跟进 {len(todo)} ｜ ✅ 已跟过(仅需修表) {len(done)}\n")
    print("=== 🔴 今天真正需要跟进（按应跟日分组）===")
    by = defaultdict(list)
    for x in todo:
        by[x["nxt"]].append(x)
    for k in sorted(by):
        print(f"\n--- 应跟 {k}（{len(by[k])} 家）---")
        for x in sorted(by[k], key=lambda y: y["num"]):
            print(f"  #{x['num']:<5}{x['name'][:30]:<32}{x['region'][:11]:<13}"
                  f"次数={x['cnt']}  实测最近={x['sent_last'] or '（无记录）'}")
    print("\n=== ✅ 已跟过、只需推进表格 ===")
    for x in sorted(done, key=lambda y: y["num"]):
        print(f"  #{x['num']:<5}{x['name'][:30]:<32}应跟 {x['nxt']}  实测最近={x['sent_last']}（Re:{x['sent_re']}）")

    OUT.write_text(json.dumps(
        {"todo": todo, "done": done,
         "generated": datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M")},
        ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n已写 {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
